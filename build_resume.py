import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml

doc = docx.Document()

# Page margins (0.5 in top/bottom, 0.6 in left/right)
for section in doc.sections:
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.6)
    section.right_margin = Inches(0.6)

# Colors matching portfolio theme
NAVY = RGBColor(16, 44, 87)
BLUE = RGBColor(15, 76, 129)
DARK = RGBColor(20, 30, 40)
GREY = RGBColor(80, 90, 100)

def add_section_heading(title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(title.upper())
    run.font.name = 'Arial'
    run.font.size = Pt(10.5)
    run.font.bold = True
    run.font.color.rgb = NAVY
    
    pBdr = parse_xml(r'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:bottom w:val="single" w:sz="12" w:space="2" w:color="0F4C81"/></w:pBdr>')
    p._p.get_or_add_pPr().append(pBdr)

# Header Table with 2 columns: Text on left, Photo on right
table = doc.add_table(rows=1, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False

col_widths = [Inches(5.6), Inches(1.5)]
for row in table.rows:
    for idx, width in enumerate(col_widths):
        row.cells[idx].width = width

cell_left = table.cell(0, 0)
p_name = cell_left.paragraphs[0]
p_name.paragraph_format.space_after = Pt(1)
run_name = p_name.add_run("GOKULARASAN")
run_name.font.name = "Arial"
run_name.font.size = Pt(20)
run_name.font.bold = True
run_name.font.color.rgb = NAVY

p_sub = cell_left.add_paragraph()
p_sub.paragraph_format.space_after = Pt(3)
run_sub = p_sub.add_run("Electrical & Electronics Engineering Student")
run_sub.font.name = "Arial"
run_sub.font.size = Pt(11)
run_sub.font.italic = True
run_sub.font.color.rgb = GREY

p_contact = cell_left.add_paragraph()
p_contact.paragraph_format.space_after = Pt(0)
r_email = p_contact.add_run("gokularasan.28eee@gmail.com  |  ")
r_email.font.size = Pt(9.5)
r_email.font.color.rgb = DARK

r_li = p_contact.add_run("LinkedIn")
r_li.font.size = Pt(9.5)
r_li.font.color.rgb = BLUE
r_li.font.underline = True

r_mid = p_contact.add_run("  |  ")
r_mid.font.size = Pt(9.5)

r_gh = p_contact.add_run("GitHub")
r_gh.font.size = Pt(9.5)
r_gh.font.color.rgb = BLUE
r_gh.font.underline = True

p_loc = cell_left.add_paragraph()
p_loc.paragraph_format.space_after = Pt(0)
r_loc = p_loc.add_run("Coimbatore, Tamil Nadu")
r_loc.font.size = Pt(9.5)
r_loc.font.color.rgb = GREY

cell_right = table.cell(0, 1)
p_img = cell_right.paragraphs[0]
p_img.alignment = WD_ALIGN_PARAGRAPH.RIGHT
try:
    p_img.add_run().add_picture("assets/gokul_portrait.jpg", width=Inches(1.2))
except Exception as e:
    print(f"Image warning: {e}")

# EDUCATION
add_section_heading("Education")

p_edu1 = doc.add_paragraph()
p_edu1.paragraph_format.space_before = Pt(3)
p_edu1.paragraph_format.space_after = Pt(1)
r1 = p_edu1.add_run("Kumaraguru College of Technology")
r1.font.bold = True
r1.font.size = Pt(10)
r1.font.color.rgb = DARK
r1_loc = p_edu1.add_run("  |  Coimbatore, Tamil Nadu")
r1_loc.font.size = Pt(9.5)
r1_loc.font.color.rgb = GREY

p_edu1_sub = doc.add_paragraph()
p_edu1_sub.paragraph_format.space_after = Pt(4)
r1_deg = p_edu1_sub.add_run("B.E., Electrical and Electronics Engineering")
r1_deg.font.italic = True
r1_deg.font.size = Pt(9.5)
r1_deg.font.color.rgb = DARK

p_edu2 = doc.add_paragraph()
p_edu2.paragraph_format.space_after = Pt(2)
r2 = p_edu2.add_run("Sree Gokulam Matric Higher Secondary School")
r2.font.bold = True
r2.font.size = Pt(9.5)
r2.font.color.rgb = DARK
r2_det = p_edu2.add_run("  |  Grade 9 – 12 | 12th: 93.83% | 11th: 90.60%")
r2_det.font.size = Pt(9.5)
r2_det.font.color.rgb = GREY

p_edu3 = doc.add_paragraph()
p_edu3.paragraph_format.space_after = Pt(4)
r3 = p_edu3.add_run("Shalom Convent Matriculation Higher Secondary School")
r3.font.bold = True
r3.font.size = Pt(9.5)
r3.font.color.rgb = DARK
r3_det = p_edu3.add_run("  |  LKG – Grade 8 | 10th: 96.40%")
r3_det.font.size = Pt(9.5)
r3_det.font.color.rgb = GREY

# INTERNSHIPS
add_section_heading("Internships")

p_int1 = doc.add_paragraph()
p_int1.paragraph_format.space_before = Pt(3)
p_int1.paragraph_format.space_after = Pt(1)
r_in1 = p_int1.add_run("Mara Signex Solar Experts — Solar Energy & EPC Operations")
r_in1.font.bold = True
r_in1.font.size = Pt(9.5)
r_in1.font.color.rgb = DARK
r_in1_d = p_int1.add_run("  |  June 2026 | 150 hrs")
r_in1_d.font.size = Pt(9)
r_in1_d.font.color.rgb = GREY

bullets_int1 = [
    "Gained hands-on exposure to On-Grid, Off-Grid, and Hybrid solar systems and their real-world applications.",
    "Researched and compared solar inverters from Sungrow, Growatt, Deye, and ABB (FIMER).",
    "Studied VFDs, harmonics, energy losses, and harmonic-reduction techniques; visited live solar plants.",
    "Prepared technical presentations and research material on solar technologies and EPC operations."
]
for b in bullets_int1:
    bp = doc.add_paragraph(style="List Bullet")
    bp.paragraph_format.space_after = Pt(1)
    br = bp.add_run(b)
    br.font.size = Pt(9)
    br.font.color.rgb = DARK

p_int2 = doc.add_paragraph()
p_int2.paragraph_format.space_before = Pt(3)
p_int2.paragraph_format.space_after = Pt(1)
r_in2 = p_int2.add_run("Electromotive Power Drives Pvt. Ltd. — Electrical Motor Manufacturing")
r_in2.font.bold = True
r_in2.font.size = Pt(9.5)
r_in2.font.color.rgb = DARK
r_in2_d = p_int2.add_run("  |  May 2025 | 7 days, 45 hrs")
r_in2_d.font.size = Pt(9)
r_in2_d.font.color.rgb = GREY

bullets_int2 = [
    "Gained firsthand exposure to electrical motor manufacturing, ISO standards, 5S, and TQM practices.",
    "Observed practical coil-winding techniques and studied key components of electrical motors.",
    "Connected theoretical electrical-machine concepts with real manufacturing processes and quality control."
]
for b in bullets_int2:
    bp = doc.add_paragraph(style="List Bullet")
    bp.paragraph_format.space_after = Pt(1)
    br = bp.add_run(b)
    br.font.size = Pt(9)
    br.font.color.rgb = DARK

# PROJECTS
add_section_heading("Projects")

p_pr1 = doc.add_paragraph()
p_pr1.paragraph_format.space_before = Pt(3)
p_pr1.paragraph_format.space_after = Pt(1)
r_pr1 = p_pr1.add_run("Real-Time Classroom Energy Monitoring System")
r_pr1.font.bold = True
r_pr1.font.size = Pt(9.5)
r_pr1.font.color.rgb = DARK
r_pr1_t = p_pr1.add_run("  |  ESP32, Current Sensor, IoT, Embedded Systems")
r_pr1_t.font.size = Pt(9)
r_pr1_t.font.color.rgb = GREY

bullets_pr1 = [
    "Built an IoT-based dashboard to track real-time current, power, energy usage, and estimated electricity cost.",
    "Enabled historical data analysis to identify peak consumption periods and reduce unnecessary energy usage.",
    "Added over-current/abnormal-consumption alerts, with scope for scaling across multiple classrooms."
]
for b in bullets_pr1:
    bp = doc.add_paragraph(style="List Bullet")
    bp.paragraph_format.space_after = Pt(1)
    br = bp.add_run(b)
    br.font.size = Pt(9)
    br.font.color.rgb = DARK

p_pr2 = doc.add_paragraph()
p_pr2.paragraph_format.space_before = Pt(3)
p_pr2.paragraph_format.space_after = Pt(1)
r_pr2 = p_pr2.add_run("Smart Cooking Oil Impurity Detection System")
r_pr2.font.bold = True
r_pr2.font.size = Pt(9.5)
r_pr2.font.color.rgb = DARK
r_pr2_t = p_pr2.add_run("  |  ESP32, pH/Moisture/Temp Sensors, IoT, Data Logging")
r_pr2_t.font.size = Pt(9)
r_pr2_t.font.color.rgb = GREY

bullets_pr2 = [
    "Designed a low-cost sensor system combining moisture, pH, and temperature sensing to flag oil contamination.",
    "Logged and compared fresh vs. used oil data with a threshold-based warning system on an IoT dashboard.",
    "Explored advanced parameters (turbidity, optical color analysis, conductivity) and ML-based classification."
]
for b in bullets_pr2:
    bp = doc.add_paragraph(style="List Bullet")
    bp.paragraph_format.space_after = Pt(1)
    br = bp.add_run(b)
    br.font.size = Pt(9)
    br.font.color.rgb = DARK

# HACKATHONS & EVENTS
add_section_heading("Hackathons & Events")

events = [
    ("Kumaraguru Livestock Hackathon 2026", "Management Team: coordinated logistics, teams, and on-ground operations for a 3-day innovation event (Jan 2026)."),
    ("U & Me Hackathon, CIT", "Participant: co-built an IoT-based Water Quality Monitoring System using sensors and embedded systems (Aug 2025)."),
    ("IGBC Poster & Model Making Competition", "Student Coordinator: managed event promotion, registrations, and on-day coordination (Oct 2025)."),
    ("Student Conclave for Climate Action (SCCA'25)", "Organizer: supported planning and execution of a student-led sustainability conclave (2025).")
]
for ev_title, ev_desc in events:
    bp = doc.add_paragraph(style="List Bullet")
    bp.paragraph_format.space_after = Pt(1)
    br1 = bp.add_run(f"{ev_title} — ")
    br1.font.bold = True
    br1.font.size = Pt(9)
    br1.font.color.rgb = DARK
    br2 = bp.add_run(ev_desc)
    br2.font.size = Pt(9)
    br2.font.color.rgb = DARK

# INDUSTRIAL EXPOSURE
add_section_heading("Industrial Exposure")

exposures = [
    ("INTEC 2026 (Intl. Machine Tools & Industrial Trade Fair, Coimbatore)", "Engaged with engineers and industry professionals on industrial machines and manufacturing systems (June 2026)."),
    ("Ambertex Universal Spinning Division", "Studied the raw-cotton-to-yarn spinning process, automation, quality control, and a ~1,500-panel rooftop solar installation (June 2026).")
]
for ex_title, ex_desc in exposures:
    bp = doc.add_paragraph(style="List Bullet")
    bp.paragraph_format.space_after = Pt(1)
    br1 = bp.add_run(f"{ex_title} — ")
    br1.font.bold = True
    br1.font.size = Pt(9)
    br1.font.color.rgb = DARK
    br2 = bp.add_run(ex_desc)
    br2.font.size = Pt(9)
    br2.font.color.rgb = DARK

# SKILLS
add_section_heading("Skills")

skills_list = [
    ("Technical", "ESP32/Embedded Systems, IoT Dashboards, Sensor Integration, Electrical Machines, Solar (On/Off-Grid, Hybrid), VFDs & Harmonics"),
    ("Tools & Concepts", "ISO Standards, 5S, TQM, Data Logging & Analysis, Technical Research & Presentation"),
    ("Soft Skills", "Event Management, Team Coordination, Leadership, Technical Communication, Problem-Solving")
]
for sk_cat, sk_items in skills_list:
    bp = doc.add_paragraph(style="List Bullet")
    bp.paragraph_format.space_after = Pt(1)
    br1 = bp.add_run(f"{sk_cat}: ")
    br1.font.bold = True
    br1.font.size = Pt(9)
    br1.font.color.rgb = DARK
    br2 = bp.add_run(sk_items)
    br2.font.size = Pt(9)
    br2.font.color.rgb = DARK

doc.save("assets/Gokularasan_Resume.docx")
print("SUCCESS: assets/Gokularasan_Resume.docx created successfully!")
